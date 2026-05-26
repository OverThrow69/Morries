from datetime import datetime
from pathlib import Path
import subprocess
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl is required to generate the weekly Excel report.")
    print("py -m pip install openpyxl")
    sys.exit(1)


PROJECT_ROOT = Path(__file__).resolve().parent.parent
PYGAME_DIR = PROJECT_ROOT / "pygame_version"
REPORTS_DIR = PYGAME_DIR / "reports"
OUTPUT_DATE_FORMAT = "%Y-%m-%d"
OUTPUT_TIME_FORMAT = "%Y-%m-%d %H:%M:%S"

PYGAME_FILES = [
    "pygame_version/README_PYGAME.md",
    "pygame_version/config.py",
    "pygame_version/main_pygame.py",
    "pygame_version/renderer.py",
    "pygame_version/reporting.py",
    "pygame_version/reports/.gitkeep",
    "pygame_version/visual_theme.py",
    "pygame_version/generate_weekly_report.py",
]

COMPILE_CHECKS = [
    ["py", "-m", "py_compile", "main.py", "game_config.py", "game_rules.py", "launch.pyw"],
    [
        "py",
        "-m",
        "py_compile",
        "pygame_version/main_pygame.py",
        "pygame_version/renderer.py",
        "pygame_version/config.py",
    ],
    ["py", "-m", "py_compile", "pygame_version/generate_weekly_report.py"],
]

VISUAL_LIMITATIONS = [
    "Pygame version is still a visual prototype shell.",
    "No real combat is implemented in the Pygame prototype.",
    "No wave spawning or wave progression is implemented in the Pygame prototype.",
    "Gold costs are visible as prototype state only; spending is not implemented.",
    "Enemy movement, projectiles, and sounds are not implemented.",
    "Units and enemies still use placeholder shapes instead of final sprite assets.",
    "Pygame state is not shared with the Tkinter gameplay implementation yet.",
]

MANUAL_TESTING_GAPS = [
    "Pygame visual checks still depend on manual window review.",
    "F9 screenshot/report workflow needs periodic manual verification.",
    "No automated screenshot comparison exists yet.",
    "No manual playtest record exists for Pygame combat because combat is not implemented.",
    "Cross-platform batch/script launch has not been manually verified outside this machine.",
]

NEXT_STEPS = [
    "Commit the weekly report generator as a safe reporting checkpoint.",
    "Generate a weekly report after each visual prototype checkpoint.",
    "Capture at least one F9 visual report before asking for visual feedback.",
    "Add Pygame gameplay only after reporting is committed and verified.",
    "Keep Tkinter gameplay and balance unchanged while Pygame remains a prototype.",
]


def run_command(command):
    completed = subprocess.run(
        command,
        cwd=PROJECT_ROOT,
        capture_output=True,
        text=True,
        shell=False,
    )
    output_parts = []
    if completed.stdout.strip():
        output_parts.append(completed.stdout.strip())
    if completed.stderr.strip():
        output_parts.append(completed.stderr.strip())
    output = "\n".join(output_parts).strip()
    return {
        "command": " ".join(command),
        "returncode": completed.returncode,
        "status": "PASS" if completed.returncode == 0 else "FAIL",
        "output": output or "(no output)",
    }


def list_report_files():
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    files = []
    for path in sorted(REPORTS_DIR.iterdir(), key=lambda item: item.name.lower()):
        if path.is_file() and path.name != ".gitkeep":
            files.append({
                "name": path.name,
                "type": path.suffix.lower() or "(none)",
                "size": path.stat().st_size,
                "modified": datetime.fromtimestamp(path.stat().st_mtime).strftime(OUTPUT_TIME_FORMAT),
            })
    return files


def get_git_branch():
    result = run_command(["git", "branch", "--show-current"])
    if result["returncode"] == 0:
        return result["output"]
    return "unknown"


def split_lines(text):
    return text.splitlines() if text else []


def add_title(sheet, title):
    sheet["A1"] = title
    sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.merge_cells("A1:F1")
    sheet.row_dimensions[1].height = 24


def write_key_values(sheet, start_row, rows):
    row = start_row
    for key, value in rows:
        sheet.cell(row=row, column=1, value=key)
        sheet.cell(row=row, column=2, value=value)
        sheet.cell(row=row, column=1).font = Font(bold=True)
        sheet.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    return row


def write_table(sheet, start_row, headers, rows):
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=start_row, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row_index, row_values in enumerate(rows, start=start_row + 1):
        for col, value in enumerate(row_values, start=1):
            cell = sheet.cell(row=row_index, column=col, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    return start_row + len(rows) + 1


def write_lines(sheet, start_row, title, lines):
    sheet.cell(row=start_row, column=1, value=title)
    sheet.cell(row=start_row, column=1).font = Font(bold=True)
    row = start_row + 1
    if not lines:
        lines = ["(none)"]
    for line in lines:
        sheet.cell(row=row, column=1, value=line)
        sheet.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        row += 1
    return row


def apply_sheet_format(sheet):
    widths = {
        "A": 30,
        "B": 48,
        "C": 28,
        "D": 22,
        "E": 22,
        "F": 22,
    }
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"


def build_workbook(data):
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    summary = workbook.create_sheet("Summary")
    add_title(summary, "Stormwall Pygame Weekly Report")
    write_key_values(summary, 3, [
        ("Generated", data["generated_at"]),
        ("Output file", str(data["output_path"])),
        ("Current git branch", data["branch"]),
        ("Unit tests", data["tests"]["status"]),
        ("Compile checks", "PASS" if all(item["returncode"] == 0 for item in data["compile_results"]) else "FAIL"),
        ("Pygame files present", f"{data['present_file_count']} of {len(PYGAME_FILES)}"),
        ("Visual report files found", len(data["report_files"])),
    ])
    write_lines(summary, 12, "Recommended Next Steps", NEXT_STEPS)

    visual = workbook.create_sheet("Visual Status")
    add_title(visual, "Visual Status")
    prototype_file_rows = [[path, "Present" if (PROJECT_ROOT / path).exists() else "Missing"] for path in PYGAME_FILES]
    next_row = write_table(visual, 3, ["Pygame Prototype File", "Status"], prototype_file_rows)
    next_row += 2
    file_rows = [[item["name"], item["type"], item["size"], item["modified"]] for item in data["report_files"]]
    if not file_rows:
        file_rows = [["No screenshot or markdown visual reports found yet.", "", "", ""]]
    next_row = write_table(visual, next_row, ["Report File", "Type", "Size (bytes)", "Modified"], file_rows)
    write_lines(visual, next_row + 2, "Known Visual Limitations", VISUAL_LIMITATIONS)

    gameplay = workbook.create_sheet("Gameplay Status")
    add_title(gameplay, "Gameplay Status")
    write_lines(gameplay, 3, "Pygame Prototype Scope", [
        "Separate Pygame visual prototype; it does not replace the Tkinter game.",
        "Current report generator is reporting-only and does not change gameplay.",
        "Tkinter gameplay and balance are intentionally untouched.",
    ])
    write_lines(gameplay, 9, "Manual Testing Gaps", MANUAL_TESTING_GAPS)

    tests = workbook.create_sheet("Tests")
    add_title(tests, "Tests")
    write_key_values(tests, 3, [
        ("Command", data["tests"]["command"]),
        ("Status", data["tests"]["status"]),
        ("Return code", data["tests"]["returncode"]),
    ])
    write_lines(tests, 8, "Unit Test Output", split_lines(data["tests"]["output"]))
    compile_start = 12 + len(split_lines(data["tests"]["output"]))
    compile_rows = [[item["command"], item["status"], item["returncode"], item["output"]] for item in data["compile_results"]]
    write_table(tests, compile_start, ["Compile Command", "Status", "Return Code", "Output"], compile_rows)

    git = workbook.create_sheet("Git Status")
    add_title(git, "Git Status")
    write_key_values(git, 3, [
        ("Current branch", data["branch"]),
        ("Git status command", data["git_status"]["command"]),
    ])
    write_lines(git, 7, "Git Status", split_lines(data["git_status"]["output"]))
    commit_start = 10 + len(split_lines(data["git_status"]["output"]))
    write_lines(git, commit_start, "Latest 5 Commits", split_lines(data["latest_commits"]["output"]))

    bugs = workbook.create_sheet("Bugs and Risks")
    add_title(bugs, "Bugs and Risks")
    write_lines(bugs, 3, "Known Risks", VISUAL_LIMITATIONS + MANUAL_TESTING_GAPS)

    next_steps = workbook.create_sheet("Next Steps")
    add_title(next_steps, "Next Steps")
    write_lines(next_steps, 3, "Recommended Next Steps", NEXT_STEPS)

    for sheet in workbook.worksheets:
        apply_sheet_format(sheet)

    return workbook


def collect_data():
    now = datetime.now()
    output_path = REPORTS_DIR / f"weekly_report_{now.strftime(OUTPUT_DATE_FORMAT)}.xlsx"
    tests = run_command(["py", "-m", "unittest", "discover", "-v"])
    compile_results = [run_command(command) for command in COMPILE_CHECKS]
    present_file_count = sum(1 for path in PYGAME_FILES if (PROJECT_ROOT / path).exists())
    return {
        "generated_at": now.strftime(OUTPUT_TIME_FORMAT),
        "output_path": output_path,
        "branch": get_git_branch(),
        "latest_commits": run_command(["git", "log", "--oneline", "-5"]),
        "git_status": run_command(["git", "status"]),
        "tests": tests,
        "compile_results": compile_results,
        "present_file_count": present_file_count,
        "report_files": list_report_files(),
    }


def main():
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    data = collect_data()
    workbook = build_workbook(data)
    workbook.save(data["output_path"])
    print(f"Weekly report saved to: {data['output_path']}")
    if data["tests"]["returncode"] != 0 or any(item["returncode"] != 0 for item in data["compile_results"]):
        print("Warning: report generated, but one or more checks failed. See the workbook for details.")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
